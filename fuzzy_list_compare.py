"""
FuzzyListCompare - Sheet-Based Pairwise Fuzzy Duplicate Comparison
==================================================================
Compares two lists in an Excel workbook for fuzzy duplicates.
Results go on a new dedicated sheet sorted by match confidence.
The best-match (secondary/target) column is colour-coded.
The original sheet is never modified.

Python port of the original VBA module (FuzzyListCompare.bas).

Usage as a library:
    from openpyxl import load_workbook
    from fuzzy_list_compare import run_fuzzy_comparison

    wb = load_workbook("data.xlsx")
    ws = wb.active
    result = run_fuzzy_comparison(wb, ws,
                                  primary_col=1, secondary_col=2,
                                  primary_header="List A",
                                  secondary_header="List B",
                                  match_threshold=0.7)
    wb.save("data.xlsx")

Usage from the command line:
    python fuzzy_list_compare.py data.xlsx --primary A --secondary B
    python fuzzy_list_compare.py data.xlsx --primary "List A" --secondary "List B" --threshold 80
"""

import argparse
import string
import sys
from collections import Counter

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

ASCII_LETTERS = set(string.ascii_letters)
ASCII_ALNUM_SPACE = set(string.ascii_letters + string.digits + " ")
NUMBERING_CHARS = set(string.digits + " .()-")

# (minimum confidence, ARGB fill colour) - checked in order
CONFIDENCE_COLOURS = [
    (0.95, "FF008000"),  # Dark green   - near-exact
    (0.80, "FF00B050"),  # Green        - strong
    (0.60, "FF92D050"),  # Yellow-green - moderate
    (0.40, "FFFFFF00"),  # Yellow       - weak
    (0.20, "FFFFA500"),  # Orange       - very weak
]
NO_MATCH_COLOUR = "FFFF0000"  # Red - no match


# ---------------------------------------------------------------------------
# clean_string - strip leading numbering, punctuation, collapse spaces, lowercase
# ---------------------------------------------------------------------------
def clean_string(s: str) -> str:
    # Step 1: Strip leading sequential numbering.
    # Skip digits, spaces, dots, parens, hyphens until the first letter
    # (or any other character that is not part of a numbering prefix).
    start = 0
    for i, ch in enumerate(s):
        if ch not in NUMBERING_CHARS:
            start = i
            break
    s = s[start:]

    # Step 2: Keep only letters, digits, and spaces
    result = "".join(ch for ch in s if ch in ASCII_ALNUM_SPACE)

    # Steps 3-5: Collapse multiple spaces, trim, lowercase
    return " ".join(result.split()).lower()


# ---------------------------------------------------------------------------
# char_bag_similarity - character-frequency similarity in [0, 1]
# ---------------------------------------------------------------------------
def char_bag_similarity(a: str, b: str) -> float:
    total_len = len(a) + len(b)
    if total_len == 0:
        return 0.0

    freq_a = Counter(a)
    freq_b = Counter(b)
    common_chars = sum(min(count, freq_b[ch]) for ch, count in freq_a.items())

    return (2.0 * common_chars) / total_len


# ---------------------------------------------------------------------------
# confidence_colour / apply_confidence_colour - cell fill by confidence score
# ---------------------------------------------------------------------------
def confidence_colour(confidence: float) -> str:
    for threshold, colour in CONFIDENCE_COLOURS:
        if confidence >= threshold:
            return colour
    return NO_MATCH_COLOUR


def apply_confidence_colour(cell, confidence: float) -> None:
    colour = confidence_colour(confidence)
    cell.fill = PatternFill(start_color=colour, end_color=colour, fill_type="solid")


# ---------------------------------------------------------------------------
# Worksheet helpers
# ---------------------------------------------------------------------------
def cell_text(value) -> str:
    """String form of a cell value ('' for empty cells)."""
    if value is None:
        return ""
    return str(value)


def last_row_in_column(ws, col: int) -> int:
    """Last row with a non-empty cell in the column (0 if the column is empty)."""
    for row in range(ws.max_row, 0, -1):
        if cell_text(ws.cell(row=row, column=col).value).strip() != "":
            return row
    return 0


def read_column_items(ws, col: int) -> list:
    """Read the data rows (2 .. last non-empty) of a column as strings."""
    last_row = last_row_in_column(ws, col)
    if last_row < 2:
        return []
    return [cell_text(ws.cell(row=r, column=col).value) for r in range(2, last_row + 1)]


def unique_sheet_name(wb, base: str) -> str:
    """Truncate to Excel's 31-char limit and de-duplicate against existing sheets."""
    name = base[:31]
    existing = {ws_name.lower() for ws_name in wb.sheetnames}
    num = 1
    while name.lower() in existing:
        num += 1
        suffix = f" {num}"
        name = base[: 31 - len(suffix)] + suffix
    return name


# ---------------------------------------------------------------------------
# run_fuzzy_comparison - main comparison logic
# ---------------------------------------------------------------------------
def run_fuzzy_comparison(wb, source_ws, primary_col: int, secondary_col: int,
                         primary_header: str, secondary_header: str,
                         match_threshold: float) -> dict:
    """
    Compare the primary column against the secondary column and write a new
    results sheet to the workbook. Returns a summary dict with keys:
    sheet_name, primary_count, secondary_count, missing_primary, missing_secondary.

    Raises ValueError if either list is empty.
    """
    # ==================================================================
    # 1. READ DATA from the source sheet
    # ==================================================================
    primary_items = read_column_items(source_ws, primary_col)
    secondary_items = read_column_items(source_ws, secondary_col)

    if not primary_items:
        raise ValueError(f"The primary list ({primary_header}) is empty.")
    if not secondary_items:
        raise ValueError(f"The secondary list ({secondary_header}) is empty.")

    primary_cleaned = [clean_string(item) for item in primary_items]
    secondary_cleaned = [clean_string(item) for item in secondary_items]

    primary_count = len(primary_items)
    secondary_count = len(secondary_items)

    # ==================================================================
    # 2. COMPUTE MATCHES
    # ==================================================================
    best_match_score = [0.0] * primary_count
    best_match_text = [""] * primary_count
    secondary_matched = [False] * secondary_count

    for i in range(primary_count):
        best_idx = -1
        if primary_cleaned[i] != "":
            for j in range(secondary_count):
                if secondary_cleaned[j] == "":
                    continue
                conf = char_bag_similarity(primary_cleaned[i], secondary_cleaned[j])
                if conf > best_match_score[i]:
                    best_match_score[i] = conf
                    best_idx = j
                    best_match_text[i] = secondary_items[j]
                    if conf >= 1.0:
                        break

        # Mark secondary item as matched if above threshold
        if best_match_score[i] >= match_threshold and best_idx >= 0:
            secondary_matched[best_idx] = True

    # ==================================================================
    # 3. SORT primary items by best match score DESCENDING (stable)
    # ==================================================================
    sort_order = sorted(range(primary_count),
                        key=lambda i: best_match_score[i], reverse=True)

    # ==================================================================
    # 4. CREATE NEW RESULTS SHEET
    # ==================================================================
    sheet_name = unique_sheet_name(wb, f"{primary_header} vs {secondary_header}")
    results_ws = wb.create_sheet(title=sheet_name, index=wb.index(source_ws) + 1)

    # ==================================================================
    # 5. WRITE RESULTS
    # ==================================================================
    bold = Font(bold=True)
    red_fill = PatternFill(start_color=NO_MATCH_COLOUR,
                           end_color=NO_MATCH_COLOUR, fill_type="solid")

    # --- Title row (merged A1:C1) ---
    threshold_pct = f"{match_threshold * 100:.0f}"
    title = (f"Comparison: {primary_header} checked against {secondary_header}"
             f" | Threshold: {threshold_pct}%")
    results_ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
    results_ws.merge_cells("A1:C1")

    # Row 2: blank spacer

    # --- Column headers (row 3) ---
    results_ws.cell(row=3, column=1, value=primary_header).font = bold
    results_ws.cell(row=3, column=2, value="Confidence").font = bold
    results_ws.cell(row=3, column=3,
                    value=f"Best Match from {secondary_header}").font = bold

    # --- Data rows (row 4 onward, sorted by confidence desc) ---
    for k, i in enumerate(sort_order):
        row = k + 4
        results_ws.cell(row=row, column=1, value=primary_items[i])
        conf_cell = results_ws.cell(row=row, column=2,
                                    value=round(best_match_score[i], 4))
        conf_cell.number_format = "0.00%"

        if best_match_score[i] > 0:
            results_ws.cell(row=row, column=3, value=best_match_text[i])

        # Colour-code the BEST MATCH (target/secondary) column
        apply_confidence_colour(results_ws.cell(row=row, column=3),
                                best_match_score[i])

    # --- "Missing from secondary" section ---
    row = primary_count + 6  # 2 blank rows after data

    missing_hdr = (f"Items in {primary_header} with no match in {secondary_header}"
                   f" (below {threshold_pct}%)")
    results_ws.cell(row=row, column=1, value=missing_hdr).font = bold
    results_ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1

    missing_primary = 0
    for i in sort_order:
        if best_match_score[i] < match_threshold:
            cell = results_ws.cell(row=row, column=1, value=primary_items[i])
            cell.fill = red_fill
            row += 1
            missing_primary += 1
    if missing_primary == 0:
        results_ws.cell(row=row, column=1, value="(none)").font = Font(italic=True)
        row += 1

    # --- "Missing from primary" section ---
    row += 2  # 2 blank rows

    missing_hdr = (f"Items in {secondary_header} with no match in {primary_header}"
                   f" (below {threshold_pct}%)")
    results_ws.cell(row=row, column=1, value=missing_hdr).font = bold
    results_ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 1

    missing_secondary = 0
    for j in range(secondary_count):
        if not secondary_matched[j]:
            cell = results_ws.cell(row=row, column=1, value=secondary_items[j])
            cell.fill = red_fill
            row += 1
            missing_secondary += 1
    if missing_secondary == 0:
        results_ws.cell(row=row, column=1, value="(none)").font = Font(italic=True)

    # ==================================================================
    # 6. FORMAT THE RESULTS SHEET
    # ==================================================================
    autofit_columns(results_ws, first_col=1, last_col=3)

    # Freeze panes at row 4 and AutoFilter on the header row
    results_ws.freeze_panes = "A4"
    results_ws.auto_filter.ref = "A3:C3"

    return {
        "sheet_name": sheet_name,
        "primary_count": primary_count,
        "secondary_count": secondary_count,
        "missing_primary": missing_primary,
        "missing_secondary": missing_secondary,
    }


def autofit_columns(ws, first_col: int, last_col: int,
                    min_width: int = 8, max_width: int = 80) -> None:
    """Approximate Excel's AutoFit by sizing columns to their longest value."""
    for col in range(first_col, last_col + 1):
        longest = 0
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            # Skip the merged title/section headers so they don't blow out col A
            if isinstance(cell.value, str) and any(
                cell.coordinate in rng for rng in ws.merged_cells.ranges
            ):
                continue
            longest = max(longest, len(cell_text(cell.value)))
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(max(longest + 2, min_width), max_width)


# ---------------------------------------------------------------------------
# Column / header resolution shared by the CLI and the GUI
# ---------------------------------------------------------------------------
def detect_list_columns(ws) -> list:
    """
    Scan row 1 for non-empty headers.
    Returns a list of (column_index, header) tuples in column order.
    """
    lists = []
    for col in range(1, ws.max_column + 1):
        header = cell_text(ws.cell(row=1, column=col).value).strip()
        if header != "":
            lists.append((col, header))
    return lists


def resolve_column(ws, spec: str):
    """
    Resolve a column given either a column letter ("A") or a header name.
    Returns (column_index, header). Raises ValueError if not found.
    """
    columns = detect_list_columns(ws)

    for col, header in columns:
        if header.lower() == spec.strip().lower():
            return col, header

    try:
        col = column_index_from_string(spec.strip().upper())
    except ValueError:
        raise ValueError(f"No column with header or letter '{spec}' found.") from None

    for c, header in columns:
        if c == col:
            return c, header
    raise ValueError(f"Column {spec.strip().upper()} has no header in row 1.")


# ---------------------------------------------------------------------------
# Command-line entry point
# ---------------------------------------------------------------------------
def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="Compare two lists in an Excel workbook for fuzzy duplicates.")
    parser.add_argument("workbook", help="Path to the .xlsx workbook")
    parser.add_argument("--sheet", help="Source sheet name (default: active sheet)")
    parser.add_argument("--primary", required=True,
                        help="Primary list: column letter or header name")
    parser.add_argument("--secondary", required=True,
                        help="Secondary list: column letter or header name")
    parser.add_argument("--threshold", type=float, default=70,
                        help="Match threshold as a percentage (default: 70)")
    parser.add_argument("--output",
                        help="Save results to this path instead of in place")
    args = parser.parse_args(argv)

    threshold = args.threshold / 100 if args.threshold > 1 else args.threshold
    if not 0 < threshold <= 1:
        parser.error("--threshold must be between 0 and 100")

    wb = load_workbook(args.workbook)
    if args.sheet:
        if args.sheet not in wb.sheetnames:
            parser.error(f"Sheet '{args.sheet}' not found in {args.workbook}")
        source_ws = wb[args.sheet]
    else:
        source_ws = wb.active

    try:
        primary_col, primary_header = resolve_column(source_ws, args.primary)
        secondary_col, secondary_header = resolve_column(source_ws, args.secondary)
        if primary_col == secondary_col:
            raise ValueError("Primary and secondary lists must be different.")

        result = run_fuzzy_comparison(wb, source_ws, primary_col, secondary_col,
                                      primary_header, secondary_header, threshold)
    except ValueError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    out_path = args.output or args.workbook
    wb.save(out_path)

    print("Comparison complete!")
    print(f"Primary list: {primary_header} ({result['primary_count']} items)")
    print(f"Secondary list: {secondary_header} ({result['secondary_count']} items)")
    print(f"Results are on sheet: {result['sheet_name']} in {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
